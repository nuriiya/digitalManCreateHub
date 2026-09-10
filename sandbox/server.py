from mcp.server.fastmcp import FastMCP
import os
import io
from pathlib import Path
import httpx
from docx import Document

mcp = FastMCP('docx-accessor')

PAPER_PROXY = os.getenv('PAPER_PROXY', '').strip()


def _load_docx_bytes(source: str, timeout: float) -> bytes:
    if source.startswith('file://'):
        source = source[7:]
    if source.startswith(('http://', 'https://')):
        proxy = PAPER_PROXY or None
        with httpx.Client(verify=False, proxy=proxy, timeout=timeout) as client:
            resp = client.get(source)
            resp.raise_for_status()
            return resp.content
    return Path(source).read_bytes()


def _extract_docx_text(data: bytes, include_tables: bool) -> str:
    document = Document(io.BytesIO(data))
    lines = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)
    if include_tables:
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip().replace(chr(10), ' ') for cell in row.cells]
                lines.append(' | '.join(cells))
    return chr(10).join(lines) if lines else '(empty document)'


@mcp.tool
def read_docx(source: str, include_tables: bool = True, timeout: float = 30.0) -> str:
    if not source or not isinstance(source, str):
        return 'Error: source must be a non-empty string path or URL'
    if timeout is None or not isinstance(timeout, (int, float)) or timeout <= 0:
        return 'Error: timeout must be a positive number'
    try:
        data = _load_docx_bytes(source, timeout)
        return _extract_docx_text(data, include_tables)
    except FileNotFoundError:
        return f'Error: file not found: {source}'
    except httpx.HTTPStatusError as exc:
        return f'Error: HTTP {exc.response.status_code} while fetching {source}'
    except Exception as exc:
        return f'Error reading docx: {exc}'


if __name__ == '__main__':
    mcp.run()